"""Build the mid-semester slide deck (16:9 .pptx) from the filled report.

    python report/build_report.py   # fills report/build/midsem_report.md (and the S8 figure)
    python report/make_slides.py    # writes report/midsem_slides.pptx

All result numbers are parsed from report/build/midsem_report.md, so the slides always match the report.
"""
import os
import re
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
REPORT = Path(os.environ.get("SLIDES_REPORT", ROOT / "report/build/midsem_report.md"))
OUT = Path(os.environ.get("SLIDES_OUT", ROOT / "report/midsem_slides.pptx"))
FIG_SYN = ROOT / "results/fig_synthetic.png"
FIG_SIM = ROOT / "results/fig_sim_m1.png"

ACCENT = RGBColor(0x1F, 0x3A, 0x5F)   # dark navy
TEXT = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x66, 0x66, 0x66)
LIGHT = RGBColor(0xEE, 0xF2, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

W, H = Inches(13.333), Inches(7.5)
MX = Inches(0.6)           # side margin
TOP = Inches(1.35)         # content top (below title)
CW = W - 2 * MX            # content width


# ----------------------------------------------------------------------------- report parsing
def md_tables(text):
    """Return every markdown table in `text` as (header, rows) in order of appearance."""
    tables, cur = [], []
    for line in text.splitlines() + [""]:
        if line.strip().startswith("|"):
            cur.append([c.strip() for c in line.strip().strip("|").split("|")])
        elif cur:
            hdr, body = cur[0], [r for r in cur[1:] if not set("".join(r)) <= set("-: ")]
            tables.append((hdr, body))
            cur = []
    return tables


def section(text, start, end):
    i = text.index(start)
    j = text.index(end, i)
    return text[i:j]


def clean(s):
    """Strip markdown/LaTeX bits for display in a table cell."""
    s = s.replace("**", "").replace("$", "")
    s = s.replace(r"\pm", "±").replace("±", "±")
    return re.sub(r"\s+", " ", s).strip()


rep = REPORT.read_text()
S4 = md_tables(section(rep, "# 3. Reproduction 1", "# 4."))[0]
S5 = md_tables(section(rep, "# 4. Reproduction 2", "# 5."))[0]
SAN = md_tables(section(rep, "# 5. MovieLens", "# 6."))[0]
_s8 = md_tables(section(rep, "# 6. Supervised", "# 7."))
S8_CLICKS, S8_CTR = _s8[0], _s8[1]

# convenient lookups for the notes (strings exactly as printed in the report)
S_VALUES = [h.replace("s = ", "").strip() for h in S8_CLICKS[0][1:]]
CLK = {r[0]: r[1:] for r in S8_CLICKS[1]}                 # agent -> ["m ± sd", ...] per s
CTR = {r[0]: r[1:] for r in S8_CTR[1]}                    # agent -> [ctr lo, ctr hi, junk lo, junk hi, len hi]
CTR_HDR = S8_CTR[0]
SYN = {r[0]: r[1:] for r in S5[1]}                        # N -> [opt, myopic, lsvi, vql]
NIP = {r[0]: r[1:] for r in S4[1]}


def mean_of(cell):
    return cell.split("±")[0].strip()


SUP, MNLQ, RAT = "Supervised (CE)", "Offline RL (MNL-Q, ours)", "Highest estimated rating (history model)"
DQN, CQL = "Online RL (DQN re-ranking top-20)", "Offline RL (CQL)"
SHORT = {RAT: "Est.-rating heuristic", "Online RL (DQN re-ranking top-20)": "Online RL (DQN, top-20)",
         "SQN-style (CE + TD aux)": "SQN-style (CE+TD)", "Offline RL (MNL-Q, ours)": "Offline RL (MNL-Q, ours)"}


# ----------------------------------------------------------------------------- drawing helpers
prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]


def _font(run, size, bold=False, color=TEXT, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def textbox(slide, x, y, w, h, text, size=18, bold=False, color=TEXT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            italic=False):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        _rich(p, line, size, bold, color, italic)
    return tb


def _rich(p, line, size, bold=False, color=TEXT, italic=False):
    """Add `line` to paragraph p; **...** segments are bold accent."""
    for k, seg in enumerate(re.split(r"\*\*", line)):
        if not seg:
            continue
        r = p.add_run()
        r.text = seg
        _font(r, size, bold or k % 2 == 1, ACCENT if (k % 2 == 1 and color == TEXT) else color, italic)


def bullets(slide, x, y, w, h, items, size=18, gap=8):
    """items: list of str or (str, level)."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        txt, lvl = (it, 0) if isinstance(it, str) else it
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        sz = size if lvl == 0 else size - 2
        mark = "■  " if lvl == 0 else "–  "
        r = p.add_run()
        r.text = ("      " * lvl) + mark
        _font(r, sz - (6 if lvl == 0 else 0), color=ACCENT if lvl == 0 else MUTED)
        _rich(p, txt, sz)
    return tb


def new_slide(title, notes=None, number=True):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = WHITE
    textbox(s, MX, Inches(0.35), CW, Inches(0.8), title, size=32, bold=True, color=ACCENT, anchor=MSO_ANCHOR.MIDDLE)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, MX, Inches(1.13), Inches(1.2), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    if number:
        textbox(s, W - MX - Inches(2.5), H - Inches(0.45), Inches(2.5), Inches(0.3),
                f"Samyak Bansal  |  {len(prs.slides)}", size=11, color=MUTED, align=PP_ALIGN.RIGHT)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def table(slide, x, y, w, header, rows, size=13, col_w=None, row_h=0.36, bold_rows=(), first_col_left=True):
    nr, nc = len(rows) + 1, len(header)
    shp = slide.shapes.add_table(nr, nc, x, y, w, Inches(row_h * nr))
    t = shp.table
    if col_w:
        tot = sum(col_w)
        for j, cw in enumerate(col_w):
            t.columns[j].width = Emu(int(w * cw / tot))
    for i in range(nr):
        t.rows[i].height = Inches(row_h)
        vals = header if i == 0 else rows[i - 1]
        for j in range(nc):
            c = t.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(0.02)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = ACCENT if i == 0 else (LIGHT if i % 2 == 0 else WHITE)
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if (j == 0 and first_col_left) else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = clean(vals[j])
            _font(r, size, bold=(i == 0 or (i - 1) in bold_rows), color=WHITE if i == 0 else TEXT)
    return shp


def image(slide, path, x, y, max_w, max_h, center=True):
    iw, ih = Image.open(path).size
    sc = min(max_w / iw, max_h / ih)
    w, h = int(iw * sc), int(ih * sc)
    if center:
        x, y = x + (max_w - w) // 2, y + (max_h - h) // 2
    return slide.shapes.add_picture(str(path), x, y, Emu(w), Emu(h))


def box(slide, x, y, w, h, title, body, size=16, fill=LIGHT):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shp.adjustments[0] = 0.08
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = ACCENT
    shp.line.width = Pt(1.25)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Inches(0.15)
    tf.margin_top = Inches(0.1)
    p = tf.paragraphs[0]
    _rich(p, title, size + 1, bold=True, color=ACCENT)
    for line in (body if isinstance(body, list) else [body]):
        q = tf.add_paragraph()
        q.space_before = Pt(5)
        _rich(q, line, size)
    return shp


def callout(slide, x, y, w, h, text, size=17):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = ACCENT
    shp.line.fill.background()
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.2)
    p = tf.paragraphs[0]
    _rich(p, text, size, bold=True, color=WHITE)
    return shp


# ----------------------------------------------------------------------------- slides
# 1. Title
s = prs.slides.add_slide(BLANK)
s.background.fill.solid()
s.background.fill.fore_color.rgb = WHITE
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.35), H)
band.fill.solid()
band.fill.fore_color.rgb = ACCENT
band.line.fill.background()
textbox(s, Inches(1.0), Inches(1.7), Inches(11.3), Inches(1.8),
        "When does RL beat a supervised recommender on long-term utility?", size=40, bold=True, color=ACCENT)
textbox(s, Inches(1.0), Inches(3.55), Inches(11.3), Inches(0.6),
        "RL for Sequential Recommendation (course project topic 3.9)  |  Mid-semester review", size=22, color=TEXT)
textbox(s, Inches(1.0), Inches(4.8), Inches(11.3), Inches(1.0),
        ["Samyak Bansal", "18 October 2026"], size=20, color=MUTED)
textbox(s, Inches(1.0), Inches(6.5), Inches(11.3), Inches(0.4),
        "Code: github.com/samyakbansal5-eng/rl-seqrec", size=14, color=MUTED)
s.notes_slide.notes_text_frame.text = (
    "Introduce yourself and the project: topic 3.9, RL for sequential recommendation. "
    "The one-line question is when an RL recommender actually beats a supervised next-click recommender on long-term "
    "utility, and what changes when we show a slate of several items instead of one. "
    "Every number in this talk is produced by a script in the repository; the README lists the commands.")

# 2. Question
s = new_slide("The question: a click now, or engagement later?", notes=(
    "A supervised sequential recommender is trained to predict the next click, so it optimises short-term "
    "engagement, i.e. click-through rate. An RL recommender maximises a discounted sum of future rewards, so in "
    "principle it can give up a click now for more engagement later. Papers report RL gains, but often on one-step "
    "metrics, and it is unclear where the gains come from. "
    "The topic asks when the RL recommender actually wins on long-term utility, and suggests the extension to "
    "re-ranking, where K items are shown at once. Our answer strategy: put both in one model so the comparison is exact."))
bullets(s, MX, TOP + Inches(0.2), Inches(7.2), Inches(5), [
    "**Supervised** next-click model: maximises **CTR now** (short-term)",
    "**RL** recommender: maximises discounted **future** clicks (long-term)",
    "RL can trade a click now for a longer, happier session",
    "Topic 3.9 asks: **when does RL actually win on long-term utility?**",
    "Extension: **re-ranking**, showing a slate of K items at once",
], size=20, gap=14)
box(s, Inches(8.3), TOP + Inches(0.3), Inches(4.4), Inches(1.6), "Short-term metric",
    ["First-slate CTR = 1 - P(no click)"], size=17)
box(s, Inches(8.3), TOP + Inches(2.2), Inches(4.4), Inches(1.6), "Long-term metric",
    ["Clicks per session (return, γ = 1) and session length"], size=17)
callout(s, MX, Inches(6.0), CW, Inches(0.7),
        "Approach: put both recommenders inside one model, so the comparison is precise rather than anecdotal", size=18)

# 3. Model
s = new_slide("One model for both: combinatorial MDP with MNL choice", notes=(
    "We use the episodic combinatorial MDP with multinomial-logit choice from Lee and Oh, ICML 2025. "
    "At each step the recommender shows a slate of at most K items; the user clicks one item or the outside option "
    "'no click' with MNL probabilities, gets reward 1 per click, and the next state depends on what was picked. "
    "Only item-level values Q-bar need to be learned; the slate value is the choice-probability-weighted average. "
    "The best slate is a linear-fractional program, solved as an LP after Charnes-Cooper or by bisection, both "
    "verified against brute force on 300 random instances. "
    "Lemma 1: with unit click rewards the best immediate slate is the top-K by utility, and a cross-entropy "
    "next-click model recovers utilities up to a constant, so supervised = the gamma = 0 policy. "
    "That turns 'supervised vs RL' into 'gamma = 0 vs gamma > 0' in one model."))
bullets(s, MX, TOP + Inches(0.15), Inches(6.4), Inches(4.6), [
    "State s: user, recent clicks, step. Action: slate A of ≤ K items",
    "User picks a ∈ A or **no click a₀** with MNL probability",
    "Reward 1 per click; next state depends on **what was picked**",
    "K = 1: sequential recommendation; K > 1: re-ranking",
    "Best slate = linear-fractional program (LP via Charnes-Cooper, or bisection; checked vs brute force)",
], size=18, gap=12)
box(s, Inches(7.3), TOP + Inches(0.15), Inches(5.4), Inches(2.45), "Choice model and slate value", [
    "P(a | s, A) = e^v(s,a) / (1 + Σ_{a'∈A} e^v(s,a'))",
    "Q̄(s,a) = r(s,a) + γ E[V(s')]",
    "Q(s,A) = Σ_{a ∈ A ∪ {a₀}} P(a | s, A) · Q̄(s,a)",
], size=17)
box(s, Inches(7.3), TOP + Inches(2.85), Inches(5.4), Inches(1.85), "Lemma 1: supervised = myopic (γ = 0)", [
    "Click prob. W/(1+W) grows in W = Σ e^v, so top-K by utility is myopic-optimal; a CE next-click model "
    "recovers utilities up to a constant",
], size=16)
callout(s, MX, Inches(6.05), CW, Inches(0.65), "\"Supervised vs RL\" becomes \"γ = 0 vs γ > 0\" in the same model",
        size=18)

# 4. Propositions
s = new_slide("When can RL help? Two propositions", notes=(
    "Proposition 1: for any n there is an MDP where the optimal long-term return is more than n times the myopic "
    "return. Construction: a junk item that is more clickable than a good item but makes the user leave; myopic always "
    "shows junk and gets at most one click, while showing the good item collects H times c_g clicks. "
    "This is the dual of Labarca, Parra and Toro Icarte's Theorem 4.3, which shows an RL-optimal policy can be "
    "arbitrarily worse at next-item prediction. So neither objective dominates: the metric and the dynamics decide. "
    "Proposition 2: if the transition does not depend on the chosen item, including no-click, the future term is a "
    "constant c(s), choice probabilities sum to one, so Q(s,A) = immediate reward + c(s) and myopic is optimal for every "
    "gamma. Corollary: RL helps only if the choice changes the future. Full proofs are in formulation.tex; "
    "Section 4 there has a two-type good/junk toy model for the dependence on K (not a proposition; to be analysed for the end-semester)."))
cw3 = (CW - Inches(0.5)) / 2
box(s, MX, TOP + Inches(0.15), cw3, Inches(3.3), "Proposition 1: RL can be arbitrarily better", [
    "For any n there is an MDP with optimal long-term return > n × myopic return.",
    "**Intuition:** clickable junk makes the user leave; myopic always shows it and gets ≤ 1 click.",
    "Dual of Labarca et al. Thm 4.3 (RL-optimal can be arbitrarily worse at next-item prediction).",
], size=16)
box(s, MX + cw3 + Inches(0.5), TOP + Inches(0.15), cw3, Inches(3.3),
    "Proposition 2: no advantage without choice-dependent dynamics", [
        "If p(s' | s, a) does not depend on the chosen item a, myopic is optimal for every γ.",
        "**Intuition:** the future term is the same for every item, so Q(s,A) = immediate reward + c(s).",
    ], size=16)
callout(s, MX, TOP + Inches(3.75), CW, Inches(0.85),
        "Corollary: RL helps only if the user's choice changes their future", size=22)
textbox(s, MX, TOP + Inches(4.85), CW, Inches(0.5),
        "Axes of the regime map: strength of that dependence (s), slate size K (two-type toy model, formulation.tex Section 4), and γ.",
        size=16, color=MUTED)

# 5. Roadmap
s = new_slide("What was built: S1 to S9", notes=(
    "This is the pipeline of what exists now. S1 data: RetailRocket and MovieLens-1M preprocessing, matching the "
    "paper's RetailRocket statistics, 195,524 sessions, 70,852 items, 57,269 purchases. S2 formulation and proofs. "
    "S3 the exact MNL assortment solver, LP and bisection, tested against brute force. "
    "S4 and S5 are the two reproductions; S6 and S7 the MovieLens simulator and its oracle sanity check; S8 the main "
    "supervised-vs-RL experiment at K = 1; S9 the Streamlit demo. "
    "Everything is a script in the repo, reproducible from the README."))
steps = [("S1", "Data", "RetailRocket,\nMovieLens-1M"), ("S2", "Theory", "Lemma 1,\nProps 1-2"),
         ("S3", "Slate solver", "LP / bisection\nMNL assortment"), ("S4", "Repro 1", "Labarca et al.\nRetailRocket"),
         ("S5", "Repro 2", "Lee & Oh\nsynthetic"), ("S6", "Simulator", "MovieLens-1M\nuser model"),
         ("S7", "Sanity", "Oracle\npolicies"), ("S8", "Main exp.", "Supervised vs\nRL, K = 1"),
         ("S9", "Demo", "Streamlit\nlive session")]
n = len(steps)
gapx = Inches(0.17)
bw = (CW - gapx * (n - 1)) / n
for i, (code, name, desc) in enumerate(steps):
    x = MX + i * (bw + gapx)
    y = TOP + Inches(0.9)
    top = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, bw, Inches(0.75))
    top.fill.solid()
    top.fill.fore_color.rgb = ACCENT
    top.line.fill.background()
    tf = top.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = code
    _font(r, 20, bold=True, color=WHITE)
    textbox(s, x - Inches(0.05), y + Inches(0.85), bw + Inches(0.1), Inches(0.5), name, size=15, bold=True,
            color=ACCENT, align=PP_ALIGN.CENTER)
    textbox(s, x - Inches(0.08), y + Inches(1.3), bw + Inches(0.16), Inches(1.0), desc.split("\n"), size=12,
            color=TEXT, align=PP_ALIGN.CENTER)
line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, MX, TOP + Inches(0.6), CW, Inches(0.04))
line.fill.solid()
line.fill.fore_color.rgb = MUTED
line.line.fill.background()
textbox(s, MX, TOP + Inches(0.15), Inches(4), Inches(0.4), "foundations", size=14, color=MUTED, italic=True)
textbox(s, MX + 3 * (bw + gapx), TOP + Inches(0.15), Inches(4), Inches(0.4), "reproductions", size=14, color=MUTED,
        italic=True)
textbox(s, MX + 5 * (bw + gapx), TOP + Inches(0.15), Inches(5), Inches(0.4), "controlled experiment + demo",
        size=14, color=MUTED, italic=True)
bullets(s, MX, TOP + Inches(3.55), CW, Inches(1.6), [
    "Status: S2-S9 done; S4 full grid on a Colab GPU with 1 seed (more seeds next)",
    "Every number in the talk comes from a script in scripts/ (README lists the commands)",
], size=18, gap=10)

# 6. Reproduction 1
hdr4, rows4 = S4
_c = {h.lower(): k for k, h in enumerate(hdr4)}
_ng = _c["buy ndcg@20"]
_pap = [k for k, h in enumerate(hdr4) if h.lower().startswith("paper")][0]
ng = {r[0]: r[_ng] for r in rows4}
n3 = {k: f"{float(v):.3f}" for k, v in ng.items()}   # 3 decimals, as in the report text
pct = lambda a, b: f"{100 * (float(ng[a]) / float(ng[b]) - 1):+.0f}%"
s = new_slide("Reproduction 1: Labarca et al. (ICML 2024) on RetailRocket", notes=(
    "Labarca, Parra and Toro Icarte claim that adding an RL loss (SQN: a Q-head trained by TD with click reward 1 and "
    "purchase reward 5) to GRU4Rec or SASRec improves next-item prediction, but that non-RL auxiliary heads (EVAL, HIST) "
    "give similar gains, so the gain is representation learning, not RL. "
    "We reproduce their protocol: time-based 80/10/10 split, hidden 64, state length 10, gamma 0.5, full-ranking "
    "HR and NDCG, model selection on validation purchase NDCG@20; our data matches the paper (195,524 sessions, 70,852 "
    "items, 57,269 purchases). The full grid, 2 backbones x 4 heads, ran on a Colab T4 GPU, 1 seed so far (paper: 5). "
    f"SASRec reproduces the paper: purchase NDCG@20 SQN {n3['SAS-SQN']} and EVAL {n3['SAS-EVAL']} vs CE {n3['SAS']}, "
    f"i.e. {pct('SAS-SQN', 'SAS')} and {pct('SAS-EVAL', 'SAS')}; EVAL has no RL in it, yet matches SQN, which is the paper's point. "
    f"GRU4Rec shows no gain: CE {n3['GRU']}, SQN {n3['GRU-SQN']}, EVAL {n3['GRU-EVAL']}; one seed, 1-6% differences, "
    "and CE stopped after 2 epochs, so this may be noise. "
    f"HIST fails in our implementation (about {n3['GRU-HIST']} and {n3['SAS-HIST']}): it regresses raw interaction counts "
    "with an unnormalised MSE that swamps the CE loss; that is our bug, to fix with a log(1+n) target, not evidence against "
    "the paper. Absolute purchase metrics are about twice the paper's (GRU CE HR@20 0.46 vs 0.23) because 67% of purchase "
    "targets equal the item viewed just before (73% in the 10-item history), so we compare relative differences."))
textbox(s, MX, TOP + Inches(0.0), CW, Inches(0.7), [
    "**Claim:** an SQN (RL) head improves next-item prediction, but non-RL heads (EVAL, HIST) do as well, so the gain is "
    "representation learning, not RL."], size=16)
cols4 = [0, _c["epochs"], _c["buy hr@20"], _ng, _c["click ndcg@20"]]
hdr_show = ["Model", "epochs", "buy HR@20", "buy NDCG@20", "click NDCG@20", "paper buy NDCG@20"]
rows_show = [[r[k] for k in cols4] + [r[_pap].split("/")[-1].strip()] for r in rows4]
table(s, MX, TOP + Inches(0.8), Inches(7.5), hdr_show, rows_show, size=12, col_w=[1.5, 1.0, 1.25, 1.35, 1.35, 1.55],
      row_h=0.36, bold_rows=[k for k, r in enumerate(rows4) if r[0] in ("SAS-SQN", "SAS-EVAL")])
textbox(s, MX, TOP + Inches(4.1), Inches(7.5), Inches(0.35),
        "Colab T4 GPU, full grid, 1 seed (paper: 5); test metrics at best validation purchase NDCG@20",
        size=12, color=MUTED, italic=True)
bullets(s, Inches(8.4), TOP + Inches(0.8), Inches(4.35), Inches(3.7), [
    f"**SASRec reproduces:** SQN {pct('SAS-SQN', 'SAS')}, EVAL {pct('SAS-EVAL', 'SAS')} purchase NDCG@20 "
    f"({n3['SAS-SQN']} / {n3['SAS-EVAL']} vs {n3['SAS']}); EVAL has no RL",
    f"**GRU4Rec: no gain** (CE {n3['GRU']} vs SQN {n3['GRU-SQN']}, EVAL {n3['GRU-EVAL']}); one seed, may be noise",
    "**HIST fails (≈0.15):** unnormalised MSE on raw counts swamps CE; our bug, fix with log(1+n)",
], size=15, gap=8)
callout(s, MX, TOP + Inches(4.55), CW, Inches(0.7),
        "Repeat items: 67% of purchase targets = the item just viewed, so absolute metrics ≈ 2× the paper's; compare relative gains",
        size=16)

# 7. Reproduction 2
hdr5, rows5 = S5
s = new_slide("Reproduction 2: Lee & Oh (ICML 2025) synthetic MDP", notes=(
    "This is the paper's synthetic 'budget' MDP: N items, 5 states, horizon 5, slates of up to 6 items, MNL choice with "
    "linear utilities, and rewards and transitions that depend on the chosen item. The optimum is computed exactly by "
    "dynamic programming with the exact assortment solver. "
    "Three agents: myopic (OFU-MNL+-style), LSVI-UCB treating each assortment as an action, and MNL-VQL, the paper's "
    "algorithm. 30,000 episodes, 5 seeds; numbers are the mean return over the last 2,000 episodes. "
    + " ".join(f"N = {k}: optimal {v[0]}, myopic {v[1]}, MNL-VQL {v[3]}." for k, v in SYN.items()) +
    " MNL-VQL is within 1% of optimal for all N; myopic is 25-37% below optimal because it ignores how the choice changes "
    f"the next state. LSVI-UCB is not in between: at N = 10 it equals myopic ({SYN['10'][2]} vs {SYN['10'][1]}, optimal {SYN['10'][0]}); it is slow because assortments grow combinatorially, so we ran it only up to N = 20. "
    "This reproduces the paper's Figure 1 qualitatively."))
table(s, MX, TOP + Inches(0.1), Inches(7.0), hdr5, rows5, size=14, col_w=[0.8, 1.6, 1.6, 1.6, 1.6], row_h=0.42)
image(s, FIG_SYN, MX, TOP + Inches(2.0), CW, Inches(2.5))
bullets(s, Inches(7.9), TOP + Inches(0.05), Inches(4.85), Inches(1.9), [
    "MNL-VQL within 1% of optimal",
    "Myopic 25-37% below optimal",
    f"LSVI-UCB = Myopic at N = 10 ({SYN['10'][2]} vs {SYN['10'][1]}); slow (combinatorial actions)",
], size=16, gap=4)
callout(s, MX, Inches(6.15), CW, Inches(0.6),
        "Takeaway: RL beats myopic when transitions depend on the choice (first clean instance)", size=17)

# 8. Simulator
s = new_slide("MovieLens-1M user simulator", notes=(
    "KuaiRec was not obtainable, so the simulator is calibrated on MovieLens-1M, as Lee and Oh do. "
    "Catalogue: 500 most-rated movies; preferences from rank-32 ALS on all 1M ratings. Junk items: the bottom 25% by "
    "mean rating, popular but poorly rated, made more clickable with kappa = 1, like clickbait. "
    "The strength s scales all long-term effects: clicking changes hidden satisfaction by s·eta·(q - q_ref), a junk click "
    "makes the user leave with probability 0.3 s, and the per-step exit probability is 0.05 + 0.4 s (1 - satisfaction); "
    "sessions are capped at 30 steps. "
    "At s = 0 nothing the user picks affects the future, so Prop. 2 says myopic is optimal: this is our control. "
    "The dynamics are assumed, not learned, so conclusions are about the mechanism."))
bullets(s, MX, TOP + Inches(0.15), Inches(7.1), Inches(4.6), [
    "**Catalogue:** 500 most-rated movies; ALS (rank 32) gives predicted rating r̂(u,i)",
    "**Junk:** bottom 25% by mean rating, made more clickable (κ = 1), like clickbait",
    "**Choice:** MNL, v = b₀ + (r̂ − 3.5) + κ·junk − s·λ·fatigue·genres",
    "**Satisfaction:** good clicks raise it, bad/junk clicks lower it",
    "**Exit:** junk click → leave w.p. 0.3 s; each step leave w.p. 0.05 + 0.4 s (1 − σ); cap 30 steps",
], size=17, gap=12)
box(s, Inches(8.1), TOP + Inches(0.15), Inches(4.6), Inches(1.6), "One knob: strength s ∈ [0, 1]", [
    "Scales every long-term effect (satisfaction, fatigue, leaving)",
], size=16)
box(s, Inches(8.1), TOP + Inches(2.0), Inches(4.6), Inches(1.6), "s = 0 is the Prop. 2 control", [
    "Choice does not affect the future, so myopic is optimal",
], size=16)

# 9. Oracle sanity check
hs, rs = SAN
s = new_slide("Sanity check: oracle policies in the simulator", notes=(
    "Before learning anything we check the simulator with two oracles that use hidden information. The myopic oracle "
    "shows the top-K by true current utility; the quality oracle shows the best predicted rating and never junk. "
    + " ".join(f"At {clean(r[0])}: myopic {r[1]}, quality {r[2]} clicks per session." for r in rs) +
    " So at s = 0 myopic wins, as Prop. 2 says; once effects are on, the quality oracle wins, and the gap grows with K, "
    "as the two-type toy model in formulation.tex Section 4 predicts. The myopic oracle always has the higher first-slate CTR."))
table(s, MX, TOP + Inches(0.3), Inches(7.6), ["setting", "myopic oracle (clicks/session)",
                                               "quality oracle (clicks/session)"], rs,
      size=16, col_w=[2.2, 2.7, 2.7], row_h=0.55, first_col_left=True)
box(s, Inches(8.7), TOP + Inches(0.3), Inches(4.0), Inches(1.9), "Oracles (use hidden state)", [
    "**Myopic:** top-K by true current utility",
    "**Quality:** best predicted rating, never junk",
], size=16)
bullets(s, MX, TOP + Inches(3.4), CW, Inches(1.8), [
    "s = 0: myopic wins (Prop. 2). With s > 0, quality wins on long-term clicks",
    "Myopic always has the higher first-slate CTR, yet loses the session",
    "Gap grows with slate size K, as the two-type toy model predicts (formulation.tex, Sec. 4)",
], size=18, gap=8)

# 10. Agents
s = new_slide("Agents compared (K = 1)", notes=(
    "Learned agents see only what a real recommender sees: user id, the last 10 clicks, the step, and factors of a "
    "separate rating model fitted on half of each user's ratings, with per-user correlation 0.77 to the simulator's "
    "true ratings. They never see satisfaction, fatigue or the junk flag. "
    "Logs: 2 sessions per training user from a popularity behaviour policy with 30% exploration; evaluation on held-out "
    "users, 3 seeds. "
    "Supervised CE uses the MNL choice likelihood, so by Lemma 1 it targets the myopic policy. SQN-style adds a TD head "
    "but still recommends from CE. CQL is conservative offline Q-learning. MNL-Q is ours: the Lee and Oh decomposition "
    "for K = 1, Q = p·Q_c + (1 - p)·N. DQN is online, re-ranking the supervised top-20."))
rows10 = [["Supervised (CE)", "MNL choice likelihood on logs", "next click (myopic, Lemma 1)"],
          ["SQN-style", "CE + TD Q-head (auxiliary)", "recommends from CE head"],
          ["Offline RL (CQL)", "logs, conservative penalty", "discounted clicks"],
          ["Offline RL (MNL-Q, ours)", "logs; p, Q_c, N heads by TD", "discounted clicks"],
          ["Online RL (DQN)", "interacts with simulator; top-20 of CE", "discounted clicks"],
          ["Est.-rating heuristic", "history rating model only", "not learned"]]
table(s, MX, TOP + Inches(0.1), Inches(8.0), ["agent", "trained on / sees", "optimises"], rows10, size=14,
      col_w=[2.5, 3.2, 2.6], row_h=0.48)
box(s, Inches(8.95), TOP + Inches(0.1), Inches(3.75), Inches(2.4), "MNL-Q (ours, K = 1)", [
    "Q(s,a) = p(a|s)·Q_c(s,a) + (1 − p(a|s))·N(s)",
    "p: click prob.; Q_c: return after click; N: after no click",
], size=15)
box(s, Inches(8.95), TOP + Inches(2.75), Inches(3.75), Inches(2.0), "What agents see", [
    "user id, last 10 clicks, step, history-model factors. **Never** satisfaction or junk flag",
], size=15)
textbox(s, MX, TOP + Inches(3.75), Inches(8.0), Inches(1.0), [
    "Logs: 2 sessions per training user (popularity policy, 30% random). Eval: held-out 20% of users, 3 seeds."],
    size=15, color=MUTED)

# 11. Main result
hdr8, rows8 = S8_CLICKS
s = new_slide("Main result: clicks per session vs strength s", notes=(
    "This is the main experiment at K = 1, mean and std over 3 seeds on held-out users. "
    + " ".join(f"At s = {sv}: supervised {mean_of(CLK[SUP][i])}, MNL-Q {mean_of(CLK[MNLQ][i])}, "
               f"DQN {mean_of(CLK[DQN][i])}, rating heuristic {mean_of(CLK[RAT][i])}."
               for i, sv in enumerate(S_VALUES) if SUP in CLK and MNLQ in CLK) +
    " At s = 0 RL brings nothing, as Prop. 2 predicts. As s grows, the supervised model keeps recommending clickable "
    "junk and loses users early; the RL agents' relative position improves with s. The y-axis of the figure is "
    "logarithmic; the dashed lines are oracles that use hidden information."))
lab = lambda a: SHORT.get(a, a)
bold_idx = [i for i, r in enumerate(rows8) if r[0] in (SUP, MNLQ)]
table(s, MX, TOP + Inches(0.05), Inches(7.3), ["agent"] + hdr8[1:], [[lab(r[0])] + r[1:] for r in rows8],
      size=12, col_w=[2.4] + [1.25] * (len(hdr8) - 1), row_h=0.4, bold_rows=bold_idx)
image(s, FIG_SIM, Inches(7.95), TOP + Inches(0.05), Inches(4.8), Inches(4.4))
textbox(s, MX, Inches(6.35), CW, Inches(0.4),
        "Mean ± std over 3 seeds, held-out users. Figure: log y-axis; dashed = oracles (hidden information).",
        size=12, color=MUTED, italic=True)

# 12. Findings
hc, rc = S8_CTR
lo, hi = S_VALUES[0], S_VALUES[-1]
adv = [round(100 * (float(mean_of(CLK[MNLQ][k])) / float(mean_of(CLK[SUP][k])) - 1)) for k in range(len(S_VALUES))]
adv_txt = ", ".join(f"{a:+d}% at s = {sv}" for a, sv in zip(adv, S_VALUES)).replace("-", "−")
junk_red = round(100 * (1 - float(CTR[MNLQ][3]) / float(CTR[SUP][3])))
s = new_slide("Findings, including the honest negatives", notes=(
    f"First, at s = 0 offline RL brings no advantage: MNL-Q is slightly worse than supervised ({mean_of(CLK[MNLQ][0])} vs "
    f"{mean_of(CLK[SUP][0])}), as Prop. 2 predicts. The online DQN is better at s = 0 ({mean_of(CLK[DQN][0])}), but for a "
    f"short-term reason: interacting with the simulator gives it better click estimates; its first-slate CTR is also the "
    f"highest, {CTR[DQN][0]}. At s = 0 the best possible policy is myopic. "
    f"Second, for s > 0 supervised keeps showing clickable junk: at s = {hi} it is no better than popularity "
    f"({mean_of(CLK[SUP][-1])} clicks/session, length {CTR[SUP][4]}). MNL-Q gives up a little CTR ({CTR[MNLQ][1]} vs "
    f"{CTR[SUP][1]}), clicks {junk_red}% less junk ({CTR[MNLQ][3]} vs {CTR[SUP][3]}), keeps users longer ({CTR[MNLQ][4]} steps) "
    f"and its advantage over supervised grows with s: {adv_txt}. "
    "Negatives: a plain DQN over all 500 items was worse than supervised (about 1.05 clicks/session at s = 1 in a single "
    "tuning run, against 1.21 for supervised in that run) because the max over noisy Q estimates picks overestimated items; the top-20 re-ranker fixed that but "
    "only matches supervised for s > 0 after 3,000 interaction steps. "
    f"And the 'highest estimated rating' heuristic beats every learned agent for s > 0 ({mean_of(CLK[RAT][-1])} at s = {hi}) "
    f"because it avoids low-rated movies by construction; the quality oracle reaches {mean_of(CLK['Oracle quality'][-1])}, "
    "so our agents recover only part of the gain. Closing that gap is a second-half goal."))
bullets(s, MX, TOP + Inches(0.1), Inches(6.6), Inches(5), [
    "**s = 0:** offline RL adds nothing (MNL-Q −4%), as Prop. 2 predicts; online DQN best, for a short-term reason (better click estimates)",
    f"**s > 0:** MNL-Q vs supervised: {adv_txt.split(', ', 1)[1]}",
    f"At s = {hi}: **{junk_red}% less junk**, longer sessions, slightly lower CTR",
    "**Negative 1:** plain DQN over 500 items worse than supervised (≈1.05 clicks/session at s = 1 in a single tuning run, vs 1.21 for supervised in that run); top-20 DQN only matches supervised for s > 0",
    "**Negative 2:** est.-rating heuristic beats all learned agents for s > 0: RL learns the direction, not the full gain",
], size=16, gap=10)
sel = [r for r in rc if r[0] in (SUP, "SQN-style (CE + TD aux)", CQL, MNLQ, DQN, RAT)]
table(s, Inches(7.45), TOP + Inches(0.1), Inches(5.3),
      ["agent", f"CTR s={lo}", f"CTR s={hi}", f"junk s={hi}", f"length s={hi}"],
      [[lab(r[0]), r[1], r[2], r[4], r[5]] for r in sel], size=12, col_w=[2.3, 1.0, 1.0, 1.0, 1.1], row_h=0.5,
      bold_rows=[i for i, r in enumerate(sel) if r[0] in (SUP, MNLQ)])
textbox(s, Inches(7.45), TOP + Inches(3.75), Inches(5.3), Inches(0.6),
        "Short-term CTR vs junk clicks and session length (from the report)", size=12, color=MUTED, italic=True)

# 13. Demo
s = new_slide("Demo: Streamlit live session", notes=(
    "The demo is app/demo.py, run with 'streamlit run app/demo.py' after training the demo models. "
    "In the live session one simulated user meets two recommenders side by side with identical random draws, so any "
    "difference is due to the policy. Per step it shows the movie (junk flagged), click probability, whether the user "
    "clicked, satisfaction and whether the user left, plus totals for clicks, session length and first-step CTR. "
    "A slider sets the strength s: at s = 0 the two sessions look alike; at high s the supervised one ends early. "
    "A second tab shows all tables and figures from the report."))
bullets(s, MX, TOP + Inches(0.2), Inches(7.3), Inches(4.5), [
    "One simulated user, **two recommenders side by side**, identical random draws",
    "Per step: movie shown (junk flagged), click prob., click, satisfaction, left?",
    "Totals: clicks, session length, first-step CTR",
    "**Slider for strength s:** watch the supervised session end early as s grows",
    "Results tab: all tables and figures of the report",
], size=18, gap=12)
box(s, Inches(8.3), TOP + Inches(0.3), Inches(4.4), Inches(1.5), "How to run", [
    "python scripts/train_demo_models.py",
    "streamlit run app/demo.py",
], size=16)

# 14. Limitations + plan
s = new_slide("Limitations and plan for pre-compre", notes=(
    "Limitations: RetailRocket has 1 seed so far (the paper uses 5) and our HIST head has a loss-scaling bug; the "
    "simulator's long-term dynamics are assumed, not learned from data, which is why s is a knob and s = 0 a control; "
    "3 seeds only; KuaiRec replaced by MovieLens-1M; datasets from GitHub mirrors. "
    "Plan: first a regime map of the RL advantage over K in {1, 3, 5, 10}, strength s and gamma, compared to the "
    "toy-model prediction. Second, a neural MNL-Q re-ranker for K > 1 with item-level choice and value heads and the exact "
    "LP assortment solver, trained offline with a conservative penalty. Third, baselines: a PISDR-style offline RL "
    "re-ranker from Yuan et al., SQN/EVAL heads, and a diversity/quality re-ranking heuristic. "
    "Then ablations on logging policy, candidate set size and gamma, closing the gap to the rating heuristic, and the "
    "full RetailRocket results."))
half = (CW - Inches(0.5)) / 2
box(s, MX, TOP + Inches(0.15), half - Inches(0.6), Inches(3.0), "Limitations", [
    "RetailRocket: 1 seed so far; HIST loss-scaling bug (fix: log(1+n))",
    "Simulator dynamics **assumed**, not learned: claims are about the mechanism",
    "3 seeds; KuaiRec replaced by MovieLens-1M",
    "Rating heuristic still beats learned agents",
], size=16)
bullets(s, MX + half - Inches(0.1), TOP + Inches(0.15), half + Inches(0.1), Inches(4.8), [
    "**Regime map:** RL advantage over K ∈ {1,3,5,10} × s × γ vs two-type toy model (formulation.tex, Sec. 4)",
    "**Neural MNL-Q re-ranker** (K > 1): item choice + value heads, exact LP slates, conservative offline",
    "**Baselines:** PISDR-style offline RL re-ranker, SQN/EVAL heads, quality re-ranking",
    "**Ablations:** logging policy, candidate set size, γ, junk vs fatigue",
    "**Close the gap** to the rating heuristic; RetailRocket: 5 seeds, fixed HIST",
], size=17, gap=10)

# 15. Questions
s = prs.slides.add_slide(BLANK)
s.background.fill.solid()
s.background.fill.fore_color.rgb = WHITE
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.35), H)
band.fill.solid()
band.fill.fore_color.rgb = ACCENT
band.line.fill.background()
textbox(s, Inches(1.0), Inches(2.4), Inches(11.3), Inches(1.2), "Questions?", size=44, bold=True, color=ACCENT)
textbox(s, Inches(1.0), Inches(3.7), Inches(11.3), Inches(1.0),
        "RL helps only if the user's choice changes their future.", size=22, color=TEXT)
textbox(s, Inches(1.0), Inches(6.3), Inches(11.3), Inches(0.4),
        "Code and report: github.com/samyakbansal5-eng/rl-seqrec  |  proofs: report/formulation.tex", size=14,
        color=MUTED)
s.notes_slide.notes_text_frame.text = (
    "Close with the one-sentence takeaway and invite questions. Likely questions: why the rating heuristic wins "
    "(our learned agents only see 2 logged sessions per user), whether the simulator is realistic (dynamics assumed; s = 0 "
    "is the control), and why GRU4Rec shows no SQN gain on RetailRocket (one seed, 1-6% differences, CE stopped after 2 epochs).")

prs.save(OUT)
print("wrote", OUT, "with", len(prs.slides), "slides")
